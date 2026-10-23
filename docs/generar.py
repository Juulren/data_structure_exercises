from pathlib import Path
import json, re
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
D=ROOT/'docs'
m=json.loads((D/'evidencias/manifest.json').read_text(encoding='utf-8'))
figures={x['name']:x['caption'] for x in m['figuras']}
specs={}
for source in ['listas','pilas','colas']:
    for line in (ROOT/f'wwwroot/js/{source}.js').read_text(encoding='utf-8-sig').splitlines():
        match=re.match(r'\s*(\d+): \{ title:',line)
        if match:
            values=re.findall(r"(?:title|description|inputs|process|output): '([^']*)'",line)
            specs[int(match[1])]=values
steps={
1:['Seleccionar Productos disponibles y retirados y pulsar Escanear producto tres veces.','Retirar producto 1 desde su fila; comparar disponibles y retirados.','Editar otro producto para comprobar la validación numérica sin modificar los datos válidos.'],
2:['Abrir Generar números y escribir 4 como cantidad.','Pulsar Generar y clasificar; comparar la lista original con las dos listas de salida.','Repetir con el campo vacío y después con abc para comprobar ambos mensajes requeridos.'],
3:['Registrar Ana con 8.5, Luis con 6.9 y Carla con 7.','Comprobar que Ana y Carla quedan aprobadas y Luis aparece en reprobados.','Intentar registrar Lucía con 11 y verificar que la lista no cambia.'],
4:['Registrar P3 Leche 26.50, P1 Arroz 30 y P2 Aceite 45.25.','Comprobar el orden por nombre y eliminar P3 por su clave.','Verificar el nuevo total e intentar reutilizar P1 para comprobar la unicidad.'],
5:['Agregar Fresa, Fresca, Orilla, Oso, Peso y Pez.','Comprobar los grupos por F, O y P; agregar Zorro para crear el grupo Z.','Intentar agregar dos palabras y observar que no se incorpora a las listas.'],
6:['Escribir UTM y pulsar Invertir con pila.','Avanzar tres pasos para observar U, T y M apilados, con M en el tope.','Avanzar hasta el final: los Pop construyen MTU y dejan vacía la pila.'],
7:['Evaluar radar y avanzar un paso para comparar R con el primer Pop.','Evaluar UTM y comprobar que U y M son distintos.','Probar una entrada compuesta solo por signos y revisar su rechazo.'],
8:['Escribir 592 y 3784 y ejecutar la suma.','Avanzar dos pasos para observar 9 + 8 = 17, el dígito 7 y el acarreo 1.','Ir al final y leer 4376 desde la pila resultado; probar después un operando con letras.'],
9:['Apilar 2, 5, 2 y 8 en ese orden.','Reemplazar el valor viejo 2 por el nuevo 9 y revisar las dos coincidencias.','Ejecutar Pop: debe salir 8; comprobar el rechazo de una entrada no numérica.'],
10:['Registrar Ana con Depósito y revisar el modal de confirmación.','Registrar Luis con Retiro y comprobar Frente, Final y tabla de clientes.','Atender ventanilla: Ana sale primero y se informa su espera; comprobar también la cola vacía.'],
11:['Iniciar la partida y esperar la llegada automática del primer coche a los 20 segundos.','Elegir un color incorrecto y luego el solicitado para pintar únicamente el coche del frente.','Dejar acumular cinco coches y revisar el modal final con pintados y tiempo total.'],
12:['Registrar ABC-123 de Ana y XYZ-456 de Luis y revisar los enlaces del callejón.','Solicitar la salida: debe retirarse ABC-123 y generarse su ticket.','Verificar que importe = segundos × 2 y comprobar el rechazo de una salida con la cola vacía.']}
extras={2:['02_faltantes'],4:['04_antes'],6:['06_apilado'],7:['07_negativo'],8:['08_pasos'],9:['09_pop'],10:['10_ingreso','10_atencion'],11:['11_fila','11_final'],12:['12_ticket']}

doc=Document(); s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.75)
s.different_first_page_header_footer=True
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(10.5 if name=='Normal' else 19 if name=='Title' else 16 if name=='Heading 1' else 12 if name.startswith('Heading') else 9)
    st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.06
    if st._element.pPr is not None:
        for border in list(st._element.pPr.findall(qn('w:pBdr'))):st._element.pPr.remove(border)
doc.styles['Normal'].paragraph_format.widow_control=True
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Página ').font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='Documentación de listas pilas y colas'
doc.core_properties.subject='Estructura de Datos Unidad II'
doc.core_properties.author=''
doc.add_page_break()

def p(text,bold_label=None):
    para=doc.add_paragraph()
    if bold_label: para.add_run(bold_label+' ').bold=True
    para.add_run(text);return para
def page(title):
    h=doc.add_heading(title,level=1);h.paragraph_format.page_break_before=True
def img(key,maxh=3.6):
    global figure
    path=D/'evidencias'/f'{key}.png'
    if not path.exists(): raise FileNotFoundError(path)
    w,h=Image.open(path).size; width=min(7.,maxh*w/h)
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True;para.paragraph_format.space_after=Pt(4)
    shape=para.add_run().add_picture(str(path),width=Inches(width))
    shape._inline.docPr.set('descr',figures[key])
    figure+=1
    cap=doc.add_paragraph(f'Figura {figure}. {figures[key]}','Caption');cap.paragraph_format.space_after=Pt(10)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
        for col,cell in enumerate(row.cells):
            cell.width=Inches(widths[col])
            tcpr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');b.append(edge)
            tcpr.append(b)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2' if i==0 else 'FFFFFF' if i%2 else 'F5F7F9');tcpr.append(sh)
            for para in cell.paragraphs:
                para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(4)
                for run in para.runs:run.font.size=Pt(9);run.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

figure=0
doc.add_paragraph('Documentación de listas pilas y colas','Title')
doc.add_paragraph('Estructura de Datos Unidad II','Subtitle')
p('Este documento presenta el funcionamiento de los doce ejercicios del laboratorio web. Cada ejercicio incluye su problema, el análisis de entradas, proceso y salida, la secuencia de uso y capturas reales de resultados y validaciones. La aplicación permite cambiar de ejercicio sin reiniciar el servidor.')
p('Las evidencias se obtuvieron el 22 de octubre de 2026 sobre la versión actual del proyecto. Los resultados muestran listas enlazadas, pilas LIFO y colas FIFO implementadas con estructuras propias en C#. Las entradas inválidas se rechazan con mensajes claros y los datos válidos se mantienen.')
img('00_inicio',3.6)
doc.add_heading('Organización del documento',level=2)
p('Módulo 1: cinco ejercicios de listas. Módulo 2: cuatro ejercicios de pilas. Módulo 3: banco, juego de coches y estacionamiento. Cada ejercicio ocupa una sección de desarrollo y otra de evidencias de prueba.')
page('Instalación arquitectura y criterios de evaluación')
p('Para ejecutar el proyecto se requiere el SDK de .NET 8 o superior y el runtime de ASP.NET Core 8. Desde la raíz del repositorio se ejecuta dotnet run y se abre la dirección local indicada por Kestrel. El frontend está en wwwroot y consume controladores REST.')
table(['Componente','Responsabilidad'],[['Models','Entidades, nodos, estructuras genéricas y respuestas.'],['Services','Reglas de negocio y estado en memoria protegido frente a operaciones concurrentes.'],['Controllers','Rutas HTTP, validación de entradas y respuestas uniformes.'],['wwwroot','HTML, estilos Material Design 3, formularios, navegación y gráficos SVG.']],[1.5,5.5])
p('Los servicios Singleton conservan los datos mientras se ejecuta el servidor. Detenerlo elimina los registros. Cada ejercicio tiene un reinicio independiente. El juego conserva su tiempo al cambiar de vista; banco y estacionamiento usan el reloj del servidor y muestran horas de Ciudad de México.')
table(['Criterio del instrumento','Evidencia incorporada'],[['Portada 20 puntos','Primera hoja reservada para insertar la portada del equipo.'],['Diseño y legibilidad 10 puntos','Capturas de la interfaz actual con controles y resultados.'],['Descripción y análisis 10 puntos','Problema y E P S en los doce ejercicios.'],['Funcionamiento 30 puntos','Pasos de uso y figuras numeradas con explicación.'],['Pruebas 30 puntos','Entradas válidas, errores, límites y resultados observados.']],[2.5,4.5])
for n in range(1,13):
    title,desc,inputs,process,output=specs[n]
    module='Listas' if n<=5 else 'Pilas' if n<=9 else 'Colas'
    page(f'Ejercicio {n} {title.replace(":", "")}')
    p(f'Módulo de {module.lower()}. {desc}')
    p(inputs,'Entradas');p(process,'Proceso');p(output,'Salida')
    doc.add_heading('Funcionamiento paso a paso',level=2)
    for i,text in enumerate(steps[n],1):p(f'{i}. {text}')
    img(f'{n:02}_resultado',4.4)
    page(f'Pruebas y evidencias del ejercicio {n}')
    extra=extras.get(n,[])
    for key in extra:
        img(key,1.5 if key=='11_fila' else 2.15 if len(extra)>1 else 2.9)
    img(f'{n:02}_error',1.6 if n==11 else 2.0 if len(extra)>1 else 2.6)
    expected=figures[f'{n:02}_error']
    table(['Caso','Resultado comprobado'],[['Operación válida',figures[f'{n:02}_resultado']],['Entrada o estado inválido',expected]],[1.4,5.6])
    if n==2:
        data=m['datos']['numeros'];p(f"Datos de la captura: original {data['original']}; pares {data['pares']}; impares {data['impares']}. La suma de los tamaños de ambas listas coincide con cuatro.")
    if n==10:
        a=m['datos']['banco'];p(f"En la atención capturada salió el turno {a['cliente']['turno']} de {a['cliente']['nombre']}. El servidor calculó {a['segundosEspera']} segundos de espera.")
    if n==11:
        g=m['datos']['juego'];p(f"La partida documentada finalizó con {g['pintados']} coche pintado y {g['tiempoSegundos']} segundos totales. La aceleración se activa cada tres aciertos: 20, 18, 16 segundos y así sucesivamente, con un mínimo de 4 segundos; esta regla se verificó en las pruebas con reloj controlado.")
    if n==12:
        t=m['datos']['ticket'];p(f"Verificación del ticket mostrado: {t['segundos']} segundos × $2.00 = ${t['importe']} MXN. Se cobran las fracciones de segundo sin redondear el tiempo a enteros.")
page('Verificación global y conclusiones')
p('Se ejecutaron nuevamente las pruebas automatizadas de C# con dotnet test Tests/DataStructureExercises.Tests.csproj --no-build. El resultado fue 75 pruebas aprobadas, cero fallos y cero pruebas omitidas. Además, las capturas documentan operaciones reales y respuestas de validación de los doce ejercicios.')
table(['Área','Comprobaciones'],[['Listas','Extracción de nodos, clasificación, umbral de 7, orden por nombre, totales decimales y grupos dinámicos.'],['Pilas','LIFO, Peek, Unicode, palíndromos, acarreos de hasta 1000 dígitos y reemplazo sin alterar el orden.'],['Colas','FIFO, sobreflujo y subflujo, reutilización, enlaces recíprocos, tiempos y cobro exactos, concurrencia y juego con reloj controlado.']],[1.2,5.8])
p('Las estructuras mantienen el orden de sus operaciones y separan la lógica del dominio de la presentación. La lista permite clasificar y reorganizar información; la pila invierte el orden de extracción; la cola atiende según llegada. Los mensajes Hay datos faltantes e Introducir sólo números se observan en las evidencias, junto con los errores específicos de cada ejercicio.')
doc.add_heading('Referencias del trabajo',level=2)
p('Instrumento de evaluación Practica_U2_ED032026.pdf, páginas 8 y 9, criterios de entrega y estructura del documento. Asignatura Estructura de Datos, Unidad II, octubre de 2026.')
p('Repositorio del proyecto: https://github.com/Juulren/data_structure_exercises')
p('Documentación y capturas: 22 de octubre de 2026. La memoria del servidor se utiliza con fines didácticos; los registros no se conservan tras detener la aplicación.')
for border in list(doc.styles.element.xpath('.//w:pBdr')):border.getparent().remove(border)
doc.save(D/'Documentacion_Unidad_II.docx')
print(f'DOCX creado con {figure} figuras incrustadas')
